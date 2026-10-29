import io

import docx
import pymupdf

from app.ingest.parse import parse
from app.ingest.structure import build_tree, chunk_section

MD = """# Technical Manual

Intro text for the whole manual.

## Engine

The engine is a 2.0 L turbo.

### Warranty

Coverage: 5 years or 100,000 km.

## Battery

### Warranty

Coverage: 8 years or 160,000 km.

| Part | Code |
|------|------|
| Cell | B-12 |
"""


def tree_of(data: bytes, name: str):
    doc = parse(data, name)
    root, flat, conf = build_tree(doc, "doc-1")
    return doc, root, flat, conf


def headings(flat):
    return [(" > ".join(n.path), n.level) for n in flat[1:]]


def test_markdown_tree_and_title():
    doc, root, flat, conf = tree_of(MD.encode(), "manual.md")
    assert doc.title == "Technical Manual"
    assert headings(flat) == [("Engine", 1), ("Engine > Warranty", 2), ("Battery", 1), ("Battery > Warranty", 2)]
    assert root.content == "Intro text for the whole manual."
    assert flat[4].blocks[-1].kind == "table"
    assert conf > 0.8


def test_ids_are_stable_and_distinct_for_same_heading():
    _, _, a, _ = tree_of(MD.encode(), "manual.md")
    _, _, b, _ = tree_of(MD.encode(), "manual.md")
    assert [n.id for n in a] == [n.id for n in b]
    assert a[2].id != a[4].id  # both are "Warranty", different parents
    _, _, c, _ = tree_of(MD.encode(), "manual.md")
    assert build_tree(parse(MD.encode(), "m.md"), "doc-2")[1][2].id != c[2].id  # scoped by document


def test_chunks_respect_sections_and_keep_offsets():
    doc, root, flat, _ = tree_of(MD.encode(), "manual.md")
    full = "\n\n".join(b.text for b in doc.blocks)
    for n in flat:
        for c in chunk_section(n, 350):
            assert c.section_id == n.id
            assert full[c.char_start : c.char_end] == c.text


def test_oversized_text_splits_under_budget():
    body = " ".join(f"Sentence number {i} is here." for i in range(400))
    doc, root, flat, _ = tree_of(f"# A\n\n## B\n\n{body}\n".encode(), "x.md")
    chunks = chunk_section(flat[-1], 100)
    assert len(chunks) > 5 and all(c.token_count <= 100 for c in chunks)
    assert len({c.id for c in chunks}) == len(chunks)


def _pdf(outline: bool) -> bytes:
    pdf = pymupdf.open()
    for title, items in [("1 Engine", [("1.1 Warranty", "Coverage is five years.")]),
                         ("2 Battery", [("2.1 Warranty", "Coverage is eight years.")])]:
        page = pdf.new_page()
        page.insert_text((72, 30), "ACME Manual - confidential", fontsize=8)
        page.insert_text((72, 100), title, fontsize=18, fontname="helv")
        y = 140
        for sub, body in items:
            page.insert_text((72, y), sub, fontsize=14)
            page.insert_text((72, y + 24), body, fontsize=10)
            y += 60
        for i in range(8):
            page.insert_text((72, y + i * 14), f"Body line {i} with ordinary text for size statistics.", fontsize=10)
    page = pdf.new_page()
    page.insert_text((72, 30), "ACME Manual - confidential", fontsize=8)
    page.insert_text((72, 100), "Appendix text only.", fontsize=10)
    if outline:
        pdf.set_toc([[1, "1 Engine", 1], [2, "1.1 Warranty", 1], [1, "2 Battery", 2], [2, "2.1 Warranty", 2]])
    return pdf.tobytes()


def test_pdf_outline_headings_pages_and_header_removal():
    doc, root, flat, conf = tree_of(_pdf(outline=True), "m.pdf")
    assert "outline" in doc.method
    assert headings(flat) == [("1 Engine", 1), ("1 Engine > 1.1 Warranty", 2),
                              ("2 Battery", 1), ("2 Battery > 2.1 Warranty", 2)]
    assert flat[4].page_start == 2
    assert not any("confidential" in b.text for b in doc.blocks)
    assert conf > 0.9


def test_pdf_inferred_headings_without_outline():
    doc, root, flat, conf = tree_of(_pdf(outline=False), "m.pdf")
    assert "inferred" in doc.method
    assert [n.heading for n in flat[1:]] == ["1 Engine", "1.1 Warranty", "2 Battery", "2.1 Warranty"]
    assert [n.level for n in flat[1:]] == [1, 2, 1, 2]
    assert conf < 0.9


def test_docx_heading_styles_and_tables():
    d = docx.Document()
    d.add_heading("Policy", 0)
    d.add_heading("Leave", 1)
    d.add_paragraph("Staff get 30 days.")
    d.add_heading("Sick leave", 2)
    d.add_paragraph("Ten days paid.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text, t.cell(1, 0).text, t.cell(1, 1).text = "Type", "Days", "Sick", "10"
    buf = io.BytesIO()
    d.save(buf)
    doc, root, flat, _ = tree_of(buf.getvalue(), "p.docx")
    assert doc.title == "Policy"
    assert headings(flat) == [("Leave", 1), ("Leave > Sick leave", 2)]
    assert flat[2].blocks[-1].kind == "table" and "Sick | 10" in flat[2].blocks[-1].text


def test_unstructured_text_is_one_section():
    doc, root, flat, conf = tree_of(b"just some words.\n\nand more words here.", "n.txt")
    assert len(flat) == 1 and conf == 0.3 and doc.warnings
