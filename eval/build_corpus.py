"""Build the synthetic evaluation corpus and its gold question set.

    cd backend && .venv/Scripts/python ../eval/build_corpus.py

Writes eval/corpus/{4 files} and eval/dataset.json. Deterministic: the same
code produces the same files. Before writing the dataset it parses every
document with the platform's own parser and checks that each gold path
exists, so a label can never point at a heading the parser did not find.
"""

import html
import io
import json
import sys
from pathlib import Path

import docx
import pymupdf

HERE = Path(__file__).parent
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(HERE.parent / "backend"))

from corpus_content import GUIDE, HANDBOOK, MANUAL, TSBS  # noqa: E402

from app.ingest.parse import parse  # noqa: E402
from app.ingest.structure import build_tree  # noqa: E402

CORPUS = HERE / "corpus"
M, T, H, G = ("voltra-e5-owners-manual.pdf", "voltra-service-bulletins.md",
              "dealer-operations-handbook.docx", "export-shipping-guide.pdf")


# ---------------------------------------------------------------- renderers
def _html_text(text: str) -> str:
    out = []
    for para in text.split("\n\n"):
        lines = para.strip().split("\n")
        if lines[0].startswith("|"):
            rows = [[c.strip() for c in ln.strip("|").split("|")] for ln in lines if not set(ln) <= set("|-: ")]
            head, body = rows[0], rows[1:]
            out.append("<table><tr>" + "".join(f"<th>{html.escape(c)}</th>" for c in head) + "</tr>"
                       + "".join("<tr>" + "".join(f"<td>{html.escape(c)}</td>" for c in r) + "</tr>" for r in body)
                       + "</table>")
        elif para.strip():
            out.append(f"<p>{html.escape(para.strip())}</p>")
    return "".join(out)


def _tree_html(node, level: int) -> str:
    heading, text, children = node
    out = f"<h{level}>{html.escape(heading)}</h{level}>" + _html_text(text)
    return out + "".join(_tree_html(c, level + 1) for c in children)


CSS = """
body { font-family: sans-serif; font-size: 10pt; }
h1 { font-size: 22pt; } h2 { font-size: 16pt; margin-top: 14pt; } h3 { font-size: 13pt; }
h4 { font-size: 11pt; font-weight: bold; }
table { border-collapse: collapse; } td, th { border: 1px solid #888; padding: 3px; }
"""


def _recorder(heads: list, page: int):
    def rec(p):
        if p.heading and p.open_close & 1 and p.heading > 1:
            heads.append((p.heading, p.text, page))
    return rec


def render_pdf(tree, footer: str, outline: bool) -> bytes:
    """HTML -> PDF with PyMuPDF Story; optional bookmarks for h2/h3; a running footer on every page."""
    story = pymupdf.Story(html=_tree_html(tree, 1), user_css=CSS)
    buf = io.BytesIO()
    writer = pymupdf.DocumentWriter(buf)
    rect = pymupdf.paper_rect("a4")
    heads: list[tuple[int, str, int]] = []
    page, more = 0, 1
    while more:
        page += 1
        dev = writer.begin_page(rect)
        more, _ = story.place(rect + (60, 60, -60, -70))
        story.element_positions(_recorder(heads, page))
        story.draw(dev)
        writer.end_page()
    writer.close()
    doc = pymupdf.open("pdf", buf.getvalue())
    for i, pg in enumerate(doc, 1):
        pg.insert_text((60, rect.height - 35), f"{footer} - page {i} of {len(doc)}", fontsize=8)
    if outline:
        doc.set_toc([[lvl - 1, text.strip(), pg] for lvl, text, pg in heads if lvl <= 3])
    doc.set_metadata({"title": tree[0]})
    return doc.tobytes(garbage=3, deflate=True)


def render_docx(tree) -> bytes:
    d = docx.Document()
    d.core_properties.title = tree[0]
    d.add_heading(tree[0], 0)

    def walk(node, level):
        heading, text, children = node
        d.add_heading(heading, level)
        for para in filter(None, (p.strip() for p in text.split("\n\n"))):
            d.add_paragraph(para)
        for c in children:
            walk(c, level + 1)

    if tree[1]:
        d.add_paragraph(tree[1])
    for c in tree[2]:
        walk(c, 1)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def render_tsbs() -> bytes:
    out = ["# Voltra Technical Service Bulletins", "",
           "Bulletins for authorised workshops. Follow the repair exactly as written.", ""]
    for code, title, applies, symptom, cause, repair, labour in TSBS:
        out += [f"## {code}: {title}", ""]
        for h, t in [("Applies To", applies), ("Symptom", symptom), ("Cause", cause), ("Repair", repair),
                     ("Labour Time", labour)]:
            out += [f"### {h}", "", t, ""]
    return "\n".join(out).encode()


# ---------------------------------------------------------------- questions
def tsb(code: str) -> str:
    return next(f"{c}: {t}" for c, t, *_ in TSBS if c == code)


def q(qid, question, expected, gold, tags):
    return {"qid": qid, "question": question, "expected_answer": expected, "answerable": True,
            "gold": [{"doc": d, "path": p} for d, p in gold], "tags": tags}


def no(qid, question, tags=()):
    return {"qid": qid, "question": question, "answerable": False, "gold": [], "tags": list(tags)}


CH, BAT, MNT, WAR, CODES = "2 Charging", "4 High-Voltage Battery", "5 Maintenance", "6 Warranty", "7 Warning Codes"
DEST = "3 Destination Requirements"
QUESTIONS = [
    q("m01", "How long does a full charge take with an 11 kW wall box?", "About 7.5 hours.",
      [(M, [CH, "2.1 Home Charging"])], ["semantic"]),
    q("m02", "How long is the high-voltage battery warranty?",
      "8 years or 160,000 km, whichever comes first; it covers capacity falling below 70 percent.",
      [(M, [BAT, "4.3 Warranty"])], ["heading_context"]),
    q("m03", "What warranty applies to the electric motors?", "6 years or 150,000 km (drive unit warranty).",
      [(M, [WAR, "6.2 Drive Unit"])], ["heading_context"]),
    q("m04", "How long is paint covered compared with rust perforation?",
      "Paint defects 3 years; rust perforation 12 years with no distance limit.",
      [(M, [WAR, "6.3 Paint and Corrosion"])], ["semantic"]),
    q("m05", "What does warning code E-204 mean and what should I do?",
      "The charging port has overheated; stop charging and let it cool for 30 minutes.",
      [(M, [CODES, "7.1 Charging Codes"])], ["exact_id"]),
    q("m06", "What does B-415 mean?", "Insulation fault: stop safely, leave the vehicle and call roadside assistance.",
      [(M, [CODES, "7.3 Battery Codes"])], ["exact_id"]),
    q("m07", "What does D-312 indicate?", "A gear selector fault; park the car and contact service.",
      [(M, [CODES, "7.2 Drivetrain Codes"])], ["exact_id"]),
    q("m08", "What does a blue light around the charging port mean?", "Scheduled charging is waiting to start.",
      [(M, [CH, "2.4 Charging Port Light"])], ["semantic"]),
    q("m09", "What is the maximum braked trailer weight?", "1,200 kg.", [(M, ["3 Driving", "3.3 Towing"])],
      ["semantic"]),
    q("m10", "What tyre pressures are recommended?", "2.5 bar front and 2.7 bar rear (cold).",
      [(M, [MNT, "5.2 Tyres"])], ["semantic"]),
    q("m11", "How often should the brake fluid be replaced?", "Every 2 years regardless of distance.",
      [(M, [MNT, "5.3 Brake Fluid"])], ["semantic"]),
    q("m12", "I'm parking the car for two months. What charge level should I leave it at?",
      "50 to 60 percent, checking the charge every three months.", [(M, [BAT, "4.2 Storage"])], ["semantic"]),
    q("m13", "How do I get into the car if the 12V battery is dead?", "Use the mechanical key hidden in the fob.",
      [(M, [MNT, "5.5 12V Battery"])], ["semantic"]),
    q("m14", "What is the peak DC charging power and how long does 10 to 80 percent take?",
      "150 kW peak; about 28 minutes from 10 to 80 percent.", [(M, [CH, "2.2 Public DC Fast Charging"])],
      ["semantic"]),
    q("m15", "Which drive mode limits motor power, and to how much?", "Eco mode, limited to 110 kW.",
      [(M, ["3 Driving", "3.1 Drive Modes"])], ["semantic"]),
    q("m16", "Compare the basic vehicle warranty with the battery warranty.",
      "Basic: 3 years or 100,000 km. Battery: 8 years or 160,000 km.",
      [(M, [WAR, "6.1 Basic Vehicle Warranty"]), (M, [BAT, "4.3 Warranty"])], ["multi_section", "heading_context"]),
    q("m17", "How often is the car serviced, and how often should the cabin air filter be replaced?",
      "Service every 24 months or 30,000 km; cabin filter every 12 months (6 months in dusty climates).",
      [(M, [MNT, "5.1 Service Schedule"]), (M, [MNT, "5.4 Cabin Air Filter"])], ["multi_section"]),
    q("m18", "What does the warranty not cover?",
      "Tyres, wiper blades, brake pads, accident damage, and damage from non-approved accessories.",
      [(M, [WAR, "6.4 Exclusions"])], ["semantic"]),
    no("m19", "What is the E5's 0-100 km/h acceleration time?"),
    no("m20", "What does warning code E-215 mean?", ["exact_id"]),

    q("t01", "What is the repair procedure in TSB-2026-014?",
      "Fit two felt pads (part 5PS-018) under the hinge clips and replace the clips with part 5PS-019.",
      [(T, [tsb("TSB-2026-014"), "Repair"])], ["exact_id"]),
    q("t02", "What causes the problem described in TSB-2026-023?",
      "A hose clamp on the cabin heater core can loosen and weep coolant.",
      [(T, [tsb("TSB-2026-023"), "Cause"])], ["exact_id"]),
    q("t03", "Which vehicles does TSB-2026-031 apply to?", "E5 built February to June 2026.",
      [(T, [tsb("TSB-2026-031"), "Applies To"])], ["exact_id"]),
    q("t04", "What labour time is allowed for TSB-2026-007?", "0.8 hours.",
      [(T, [tsb("TSB-2026-007"), "Labour Time"])], ["exact_id"]),
    q("t05", "How is a door handle that stays retracted in freezing weather fixed?",
      "Replace the handle actuator seal with part 5DH-220 and widen the drain channel using tool T-41.",
      [(T, [tsb("TSB-2026-007"), "Repair"])], ["semantic", "heading_context"]),
    q("t06", "The centre screen keeps restarting on long drives. What is the fix?",
      "Update the head unit firmware to version 7.4.1.", [(T, [tsb("TSB-2026-019"), "Repair"])],
      ["semantic", "heading_context"]),
    q("t07", "A customer's phone key stopped working after a phone update. What should the technician do?",
      "Have the customer update the Voltra app to 5.0 or later, then delete and re-pair the phone key.",
      [(T, [tsb("TSB-2026-036"), "Repair"])], ["semantic", "heading_context"]),
    q("t08", "Why does DC charging stay below 50 kW in cold weather on some cars?",
      "Early software sets the battery heater target temperature too low.",
      [(T, [tsb("TSB-2026-011"), "Cause"])], ["semantic", "heading_context"]),
    q("t09", "What is the part number of the door handle actuator seal?", "5DH-220.",
      [(T, [tsb("TSB-2026-007"), "Repair"])], ["exact_id"]),
    q("t10", "What are the fixes for a tailgate that won't close fully and for wind noise from the driver window?",
      "Recalibrate the anti-pinch sensor with routine TG-CAL-2; refit the upper window seal with gauge T-55 "
      "and replace it (part 5WS-112) if deformed.",
      [(T, [tsb("TSB-2026-027"), "Repair"]), (T, [tsb("TSB-2026-031"), "Repair"])],
      ["multi_section", "heading_context"]),
    no("t11", "What does TSB-2026-040 cover?", ["exact_id"]),

    q("h01", "How much deposit does an export order need, and is it refundable?",
      "30 percent, non-refundable once the vehicle is allocated.", [(H, ["Export Sales", "Deposits"])],
      ["heading_context"]),
    q("h02", "What is the deposit for a domestic order and when can it be refunded?",
      "10 percent; fully refundable within 7 days, afterwards minus a 500 dirham fee.",
      [(H, ["Vehicle Sales", "Deposits"])], ["heading_context"]),
    q("h03", "How long is a trade-in valuation valid?", "14 days or 500 km, whichever comes first.",
      [(H, ["Vehicle Sales", "Trade-Ins"])], ["semantic"]),
    q("h04", "Can a customer return a new car they just bought, and on what conditions?",
      "Yes, within 5 days of delivery if driven less than 300 km and undamaged.",
      [(H, ["Returns and Cancellations", "Cooling-Off Period"])], ["semantic"]),
    q("h05", "Does the cooling-off period apply to fleet orders?", "No, fleet orders are excluded.",
      [(H, ["Returns and Cancellations", "Exceptions"])], ["semantic"]),
    q("h06", "How long are customer records kept?", "6 years after the last transaction.",
      [(H, ["Customer Data", "Retention"])], ["semantic"]),
    q("h07", "When is a customer offered a courtesy car?",
      "When a repair is expected to take more than 4 hours (and the customer has held a licence for 2+ years).",
      [(H, ["Workshop", "Courtesy Cars"])], ["semantic"]),
    q("h08", "Within how many days must warranty claims be submitted to Voltra?",
      "Within 14 days of the repair, with photos of the failed part and the diagnostic report.",
      [(H, ["Workshop", "Warranty Claims"])], ["semantic"]),
    q("h09", "What driving licence requirement applies to test drives and to courtesy cars?",
      "Both require a licence held for at least 2 years.",
      [(H, ["Vehicle Sales", "Test Drives"]), (H, ["Workshop", "Courtesy Cars"])], ["multi_section"]),
    q("h10", "Are workshops open on Fridays?", "No, workshops are closed on Friday.",
      [(H, ["Workshop", "Opening Hours"])], ["semantic"]),
    no("h11", "What commission does a sales consultant earn per car sold?"),

    q("g01", "What is the maximum vehicle age for export to Dunmark?", "3 years.",
      [(G, [DEST, "3.4 Dunmark", "3.4.1 Age Limit"])], ["heading_context"]),
    q("g02", "Can right-hand-drive cars be exported to Cestria?", "Yes, both left- and right-hand drive are accepted.",
      [(G, [DEST, "3.3 Cestria", "3.3.2 Steering"])], ["heading_context"]),
    q("g03", "What import duty does Borovia charge on electric vehicles?",
      "Electric vehicles are exempt (the standard duty is 20 percent).",
      [(G, [DEST, "3.2 Borovia", "3.2.4 Duty"])], ["heading_context"]),
    q("g04", "How far ahead should the export certificate be applied for, and what does it cost?",
      "At least 5 working days before shipping; 200 dirhams.",
      [(G, ["1 Documentation", "1.1 Export Certificate"])], ["semantic"]),
    q("g05", "Which shipping method suits a non-running vehicle?", "Container shipping.",
      [(G, ["2 Shipping Methods", "2.2 Container"])], ["semantic"]),
    q("g06", "What charge level is required for an electric vehicle shipped by air?", "At or below 30 percent.",
      [(G, ["2 Shipping Methods", "2.3 Air Freight"])], ["semantic"]),
    q("g07", "Which destinations require a certificate of origin?", "Borovia and Dunmark.",
      [(G, ["1 Documentation", "1.3 Certificate of Origin"])], ["semantic"]),
    q("g08", "Compare the inspection requirements for Arland and Borovia.",
      "Arland: pre-shipment inspection by an approved agency. Borovia: inspection on arrival, no pre-shipment check.",
      [(G, [DEST, "3.1 Arland", "3.1.3 Inspection"]), (G, [DEST, "3.2 Borovia", "3.2.3 Inspection"])],
      ["multi_section", "heading_context"]),
    q("g09", "For a RoRo export, how long from full payment to loading, and how long is transit to Arland?",
      "About 3 weeks from full payment to loading; typical transit to Arland is 12 days.",
      [(H, ["Export Sales", "Lead Times"]), (G, ["2 Shipping Methods", "2.1 RoRo"])],
      ["multi_section", "cross_document"]),
    no("g10", "What is the import duty for vehicles shipped to Estoria?"),
    no("g11", "How much does RoRo shipping to Arland cost?"),
]


def main() -> None:
    CORPUS.mkdir(exist_ok=True)
    files = {
        M: render_pdf(MANUAL, "Voltra E5 Owner's Manual", outline=True),
        T: render_tsbs(),
        H: render_docx(HANDBOOK),
        G: render_pdf(GUIDE, "Export Shipping Guide", outline=False),
    }
    paths: dict[str, set[tuple[str, ...]]] = {}
    for name, data in files.items():
        (CORPUS / name).write_bytes(data)
        doc = parse(data, name)
        _, flat, conf = build_tree(doc, "check")
        paths[name] = {tuple(n.path) for n in flat}
        print(f"{name:<36} {doc.method:<16} sections={len(flat):<3} confidence={conf}")
    missing = [(x["qid"], g["doc"], g["path"]) for x in QUESTIONS for g in x["gold"]
               if tuple(g["path"]) not in paths[g["doc"]]]
    if missing:
        for m in missing:
            print("MISSING GOLD PATH", m)
        sys.exit(1)
    dataset = {
        "name": "voltra-synthetic", "version": "v1",
        "description": "Synthetic, deliberately structure-heavy corpus (fictional EV maker and dealer). "
                       "Results describe this kind of document, not documents in general.",
        "corpus": list(files), "queries": QUESTIONS,
    }
    (HERE / "dataset.json").write_text(json.dumps(dataset, indent=2), encoding="utf-8")
    tags: dict[str, int] = {}
    for x in QUESTIONS:
        for t in x["tags"] + ([] if x["answerable"] else ["unanswerable"]):
            tags[t] = tags.get(t, 0) + 1
    print(f"{len(QUESTIONS)} questions; tags: {tags}")


if __name__ == "__main__":
    main()
